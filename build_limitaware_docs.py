#!/usr/bin/env python3
"""Build the LimitAware manuscript and supplement (DOCX + PDF) from Markdown."""
import pathlib, subprocess, sys
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

HERE = pathlib.Path(__file__).resolve().parent
ROOT = HERE if ((HERE / "pyproject.toml").exists() or (HERE / "src" / "limitaware").exists()) else HERE.parent
DIR = ROOT / "manuscript_limitaware" if (ROOT / "manuscript_limitaware").exists() else ROOT / "manuscript"
TEMPLATE = ROOT / "manuscript" / "journal_template" / "softwarex-osp-template.docx"
if not TEMPLATE.exists():
    TEMPLATE = ROOT / "manuscript" / "softwarex-osp-template.docx"
WIDTHS = {
    "manuscript": {0: [0.45, 2.05, 3.55], 1: [1.35, 2.05, 2.9],
                   2: [1.55, 0.55, 0.75, 0.75, 0.7, 2.0],
                   3: [1.0, 0.55, 0.85, 1.0, 0.75, 1.0, 0.85],
                   4: [1.75, 0.78, 0.60, 0.80, 0.78, 0.60, 0.80]},
    "supplement": {0: [1.2, 1.1, 2.4, 1.6], 1: [1.0, 0.7, 2.3, 2.3],
                   2: [2.05, 0.9, 1.75, 1.6], 3: [1.9, 0.6, 3.8],
                   4: [0.9, 0.5, 0.85, 0.55, 0.75, 0.7, 0.7, 1.35],
                   5: [1.1, 0.55, 0.8, 1.0, 0.75, 1.0, 0.85],
                   6: [1.55, 0.5, 0.6, 0.6, 0.6, 0.55, 0.6, 0.6, 0.55],
                   7: [1.4, 0.7, 1.3, 1.3, 1.6]},
}
FONT_PT = {"manuscript": {0: 8.5, 2: 8.8, 3: 8.8, 4: 8.0}, "supplement": {0: 8.2, 1: 8.5, 2: 8.5, 3: 9.0, 4: 7.8, 5: 8.2, 6: 8.5, 7: 7.6}}


def fix_tables(path, spec):
    d = docx.Document(str(path))
    for i, table in enumerate(list(d.tables)):
        widths = WIDTHS[spec].get(i, [6.3 / max(len(table.columns), 1)] * len(table.columns))
        data = [[c.text for c in row.cells] for row in table.rows]
        new = d.add_table(rows=len(data), cols=len(data[0]))
        try:
            new.style = d.styles["Table Grid"]
        except KeyError:
            pass
        pt = FONT_PT[spec].get(i, 9.5)
        for r, vals in enumerate(data):
            for c, v in enumerate(vals):
                cell = new.cell(r, c); cell.text = v
                for run in cell.paragraphs[0].runs:
                    run.font.size = docx.shared.Pt(pt)
                    if r == 0:
                        run.bold = True
        tblPr = new._tbl.tblPr
        for tag in ("w:tblW", "w:tblLayout"):
            for el in tblPr.findall(qn(tag)): tblPr.remove(el)
        w = OxmlElement("w:tblW"); w.set(qn("w:type"), "pct"); w.set(qn("w:w"), "5000"); tblPr.append(w)
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
        grid = new._tbl.find(qn("w:tblGrid"))
        if grid is not None:
            for j, col in enumerate(grid.findall(qn("w:gridCol"))):
                if j < len(widths): col.set(qn("w:w"), str(int(widths[j] * 1440)))
        for r, row in enumerate(new.rows):
            trPr = row._tr.get_or_add_trPr()
            cant = OxmlElement("w:cantSplit"); trPr.append(cant)
            if r == 0:
                hdr = OxmlElement("w:tblHeader"); trPr.append(hdr)
            for j, cell in enumerate(row.cells):
                if j < len(widths): cell.width = Inches(widths[j])
        table._tbl.addprevious(new._tbl); table._tbl.getparent().remove(table._tbl)
    # The journal reference template enables continuous line numbering. It is a
    # review aid, not part of the article layout; remove it from generated files
    # while leaving the template untouched.
    for section in d.sections:
        ln = section._sectPr.find(qn("w:lnNumType"))
        if ln is not None:
            section._sectPr.remove(ln)
    # Make reference numbering explicit so it survives DOCX/PDF conversion.
    start = None
    for i, para in enumerate(d.paragraphs):
        if para.text.strip() == "References":
            start = i
    if start is not None:
        n = 0
        for para in d.paragraphs[start + 1:]:
            if not para.text.strip():
                continue
            n += 1
            pPr = para._p.get_or_add_pPr()
            numPr = pPr.find(qn("w:numPr"))
            if numPr is not None:
                pPr.remove(numPr)
            if para.runs:
                para.runs[0].text = f"{n}. " + para.runs[0].text
            else:
                para.add_run(f"{n}. " + para.text)
    d.save(str(path))


def build(md, docx_path, spec, pdf=True):
    subprocess.run(["pandoc", md.name, "-o", str(docx_path),
                    f"--reference-doc={TEMPLATE}"], check=True, cwd=DIR)
    fix_tables(docx_path, spec)
    if pdf:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf",
                        "--outdir", str(DIR), str(docx_path)], check=True,
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    print("built", docx_path.name, "pdf" if pdf else "")


if __name__ == "__main__":
    build(DIR / "limitaware_manuscript.md", DIR / "limitaware_manuscript.docx", "manuscript")
    build(DIR / "limitaware_supplement.md", DIR / "limitaware_supplement.docx", "supplement")
